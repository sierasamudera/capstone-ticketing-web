"""Test CLI dan utilitas generator ID."""

from __future__ import annotations

import json
import os
import subprocess
import sys
from datetime import datetime

import pytest

from ticketing.cli import main as cli_main
from ticketing.seed import (
    DEFAULT_DATA_PATH,
    TicketIdGenerator,
    add_ticket,
    load_dataset,
    save_dataset,
)

import tempfile


@pytest.fixture
def tmp_dataset(tmp_path):
    """Salin dataset default ke file sementara untuk uji tulis."""
    ds = load_dataset()
    path = str(tmp_path / "tickets_test.json")
    save_dataset(ds, path)
    return path


# ---------------------------------------------------------------------------
# Generator ID (Catatan_Revisi #8: counter terpisah per prefix)
# ---------------------------------------------------------------------------


def test_id_generator_counter_terpisah():
    gen = TicketIdGenerator(["T-001", "T-002", "TK-001"])
    assert gen.next("User Report") == "T-003"
    assert gen.next("Technician Request") == "TK-002"
    assert gen.next("User Report") == "T-004"


def test_id_generator_kosong():
    gen = TicketIdGenerator()
    assert gen.next("User Report") == "T-001"
    assert gen.next("Technician Request") == "TK-001"


def test_id_generator_format():
    from ticketing.models import TicketType

    gen = TicketIdGenerator(["T-999", "TK-009"])
    assert gen.next(TicketType.USER_REPORT) == "T-1000"
    assert gen.next(TicketType.TECHNICIAN_REQUEST) == "TK-010"


# ---------------------------------------------------------------------------
# add_ticket
# ---------------------------------------------------------------------------


def test_add_ticket_user_report(tmp_dataset):
    ds = load_dataset(tmp_dataset)
    n = len(ds.tickets)
    t = add_ticket(
        ds, ticket_type="User Report", created_by="User",
        timestamp="2026-09-23 14:00", category="Network",
        detail="Tes", priority="High", duration=30,
    )
    assert len(ds.tickets) == n + 1
    assert t.id == "T-021"  # lanjut dari T-020
    assert t.weight == 3
    assert t.is_root_case is True


def test_add_ticket_technician_request(tmp_dataset):
    ds = load_dataset(tmp_dataset)
    t = add_ticket(
        ds, ticket_type="Technician Request", created_by="Admin",
        timestamp="2026-09-23 14:00", category="Network",
        detail="Bantuan", priority="Medium", duration=20,
        related_id="T-001", requested_by="T03",
    )
    assert t.id == "TK-003"  # lanjut dari TK-002
    assert t.is_root_case is False
    assert t.weighted_load == 40


def test_add_ticket_id_tidak_bentrok(tmp_dataset):
    """Catatan_Revisi #8: tidak ada ID ganda."""
    ds = load_dataset(tmp_dataset)
    ids = {t.id for t in ds.tickets}
    for i in range(5):
        t = add_ticket(
            ds, ticket_type="User Report", created_by="User",
            timestamp="2026-09-23 15:00", category="Hardware",
            detail=f"Tes {i}", priority="Low", duration=10,
        )
        assert t.id not in ids
        ids.add(t.id)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _run_cli(args, data=None):
    if data is not None:
        args = ["--data", data] + args
    return cli_main(args)


def test_cli_run_text(capsys):
    assert _run_cli(["run", "--scenario", "all"]) == 0
    out = capsys.readouterr().out
    assert "SIMULASI DISTRIBUSI BEBAN TIKET" in out
    assert "S1 Baseline" in out
    assert "S4 Redistribution" in out
    # angka kunci dari Excel
    assert "191.25" in out  # beban ideal
    assert "120" in out    # gap S1 (menit)


def test_cli_run_detail(capsys):
    assert _run_cli(["run", "--scenario", "S2", "--detail"]) == 0
    out = capsys.readouterr().out
    assert "Detail Assignment S2" in out
    assert "T-001" in out


def test_cli_run_json(capsys):
    assert _run_cli(["run", "--scenario", "S1", "--format", "json"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert "S1" in payload
    assert payload["S1"]["gap_minutes"] == 120
    assert len(payload["S1"]["assignments"]) == 22


def test_cli_run_skenario_invalid(capsys):
    with pytest.raises(SystemExit):
        _run_cli(["run", "--scenario", "S9"])


def test_cli_tickets_list(capsys):
    assert _run_cli(["tickets", "list"]) == 0
    out = capsys.readouterr().out
    assert "T-001" in out
    assert "TK-002" in out
    assert "Total: 22 tiket" in out


def test_cli_technicians(capsys):
    assert _run_cli(["technicians"]) == 0
    out = capsys.readouterr().out
    assert "T01" in out and "Teknisi A" in out


def test_cli_tickets_add_dan_simpan(tmp_dataset, capsys):
    assert _run_cli([
        "tickets", "add", "--type", "user", "--time", "2026-09-23 14:00",
        "--category", "Software", "--detail", "Tes CLI", "--priority", "Medium",
        "--duration", "25",
    ], data=tmp_dataset) == 0
    ds = load_dataset(tmp_dataset)
    assert len(ds.tickets) == 23
    assert ds.tickets[-1].id == "T-021"
    assert ds.tickets[-1].detail == "Tes CLI"


def test_cli_tickets_add_tech_request_wajib_related(tmp_dataset, capsys):
    rc = _run_cli([
        "tickets", "add", "--type", "tech", "--time", "2026-09-23 14:00",
        "--category", "Software", "--detail", "Bantuan", "--priority", "Medium",
        "--duration", "25",
    ], data=tmp_dataset)
    assert rc == 2
    ds = load_dataset(tmp_dataset)
    # tidak ada tiket bertambah
    assert len(ds.tickets) == 22


def test_cli_export_csv(tmp_path):
    out = str(tmp_path / "hasil.csv")
    assert _run_cli(["run", "--scenario", "all", "--format", "csv", "--out", out]) == 0
    assert os.path.exists(out)
    detail = str(tmp_path / "hasil_detail.csv")
    assert os.path.exists(detail)
    with open(detail, encoding="utf-8-sig") as f:
        lines = f.read().strip().split("\n")
    assert len(lines) == 23  # header + 22 tiket


def test_cli_export_xlsx(tmp_path):
    out = str(tmp_path / "hasil.xlsx")
    assert _run_cli(["run", "--scenario", "all", "--format", "xlsx", "--out", out]) == 0
    assert os.path.exists(out)
    import openpyxl

    wb = openpyxl.load_workbook(out)
    assert "Ringkasan" in wb.sheetnames
    assert any(s.startswith("Sim_S") for s in wb.sheetnames)


def test_cli_compare_excel(capsys):
    excel = r"C:\Users\SUPER USER\OneDrive - Universitas Terbuka\TUTON SEMESTER 7\Capstone Project\Tugas 1\Simulasi_Workload_Ticketing_revisi.xlsx"
    if not os.path.exists(excel):
        pytest.skip("file Excel asli tidak ditemukan")
    rc = _run_cli(["compare-excel", excel])
    out = capsys.readouterr().out
    assert rc == 0
    assert "IDENTIK dengan Excel" in out
    assert "42" in out


def test_cli_compare_excel_file_tidak_ada(capsys):
    rc = _run_cli(["compare-excel", "file_tidak_ada.xlsx"])
    assert rc == 2
    assert "tidak ditemukan" in capsys.readouterr().err


def test_cli_subprocess():
    """Uji `python -m ticketing` end-to-end lewat subprocess."""
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    r = subprocess.run(
        [sys.executable, "-m", "ticketing", "run", "--scenario", "S1"],
        capture_output=True, text=True, cwd=root,
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    assert r.returncode == 0
    assert "S1 Baseline" in r.stdout
