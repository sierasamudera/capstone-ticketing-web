import docx

doc = docx.Document()
doc.add_heading('BAB III. METODOLOGI PENELITIAN', 0)

doc.add_heading('3.1 Metode Penelitian', level=1)
doc.add_paragraph('Penelitian ini menggunakan metode Eksperimental Berbasis Simulasi Komputasi dengan pendekatan kuantitatif. Penelitian berfokus pada simulasi algoritma distribusi beban kerja pada sistem ticketing layanan IT.')
doc.add_paragraph('Pengumpulan data dilakukan melalui studi dokumen (mengkaji Simulasi_Workload_Ticketing_revisi.xlsx) dan observasi berbasis skenario (menguji 4 skenario S1-S4 pada dataset tiket).')

doc.add_heading('3.2 Proses Analisis', level=1)
doc.add_paragraph('Tahap analisis mengidentifikasi kelemahan pada aturan alokasi beban konvensional dan menentukan kriteria evaluasi performa sistem ticketing. Kelemahan utama: ketidakmerataan beban kerja (Round-Robin/S1), efektif penundaan (buktikan prioritas tinggi tertimbun), dan ketersediaan fitur redistribusi antrean (belum ada di Excel). Diketahui 4 indikator performa: Workload Gap, Jain Fairness Index, Standard Deviation, dan Average Waiting Time.')

doc.add_heading('3.3 Proses Perancangan', level=1)
doc.add_paragraph('Solusi dirancang secara modular dengan arsitektur Python Engine & Web Application. Setiap skenario memiliki logika skor tersendiri (S1: skor=0/Round-Robin, S2: durasi aktif, S3: durasi x bobot prioritas, S4: redistribusi otomatis). Rules penentu: skor terkecil, tie-breaker ID teknisi terkecil, start time = max(arrival, free_at), fallback semua busy = Earliest Estimated Finish. Struktur data: model Ticket, Technician, Parameters; generator ID otomatis TicketIdGenerator. UI: 62 halaman mockup multi-role (Admin, Pelapor, Teknisi) + Dashboard Flask.')

doc.add_heading('3.4 Implementasi', level=1)
doc.add_paragraph('Implementasi dengan tumpukan teknologi utama: 1) Bahasa Pemrograman: Python 3.12; 2) Backend Framework: Flask & WSGI; 3) Penyimpanan Data: JSON format (data/tickets.json); 4) Modul Arsitektur (ticketing/): engine.py, scenarios.py, metrics.py, cli.py, web/app.py & charts.py. Sumber daya: 63 unit test Pytest, 42 titik validasi silang terhadap Excel asli.')

doc.add_heading('3.5 Pengujian dan Evaluasi', level=1)
doc.add_paragraph('Pengujian perangkat lunak menggunakan Pytest: 63 test terbagi 5 modul (test_engine 11 test, test_cli 19 test, test_metrics 12 test, test_charts 13 test, test_seed 8 test). Validasi silang terhadap Excel menunjukkan akurasi 100% pada S1/S2/S3. Hasil evaluasi skenario: S1 gap terbesar & Jain terendah; S2 & S3 meratakan beban mendekati Jain 1.00; S4 menurunkan waiting time antrean padat.')