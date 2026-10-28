"""Command-line interface.

Pemakaian::

    python -m ticketing run --scenario all
    python -m ticketing run --scenario S2 --detail
    python -m ticketing run --scenario all --format xlsx --out hasil.xlsx
    python -m ticketing tickets list
    python -m ticketing tickets add --type user --category Network --priority High \\
        --duration 45 --detail "WiFi mati" --time "2026-09-23 14:00"
    python -m ticketing compare-excel "path/Simulasi_Workload_Ticketing_revisi.xlsx"
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import datetime

from . import __version__
from .metrics import compute_indicators
from .report import (
    assignment_table,
    export_assignments_csv,
    export_csv,
    export_xlsx,
    text_report,
)
from .scenarios import SCENARIOS
from .seed import DEFAULT_DATA_PATH, add_ticket, load_dataset, save_dataset
from .simulation import Simulation

EXCEL_PATH = r"C:\Users\SUPER USER\OneDrive - Universitas Terbuka\TUTON SEMESTER 7\Capstone Project\Tugas 1\Simulasi_Workload_Ticketing_revisi.xlsx"


def _parse_time(value: str) -> datetime:
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    raise argparse.ArgumentTypeError(
        f"format waktu tidak valid: {value!r} (pakai 'YYYY-MM-DD HH:MM')"
    )


def cmd_run(args) -> int:
    sim = Simulation.load(args.data)
    # Override parameters with external config if present (MVP uses config.yaml in project root)
    try:
        from .config import load_parameters
        sim.dataset.parameters = load_parameters()
    except Exception:
        # If config missing or broken we keep JSON defaults – harmless for MVP
        pass
    keys = list(SCENARIOS) if args.scenario in ("all", None) else [args.scenario]

    indicators = {k: sim.run(k) for k in keys}

    if args.format == "json":
        payload = {}
        for k, ind in indicators.items():
            payload[k] = {
                "name": ind.name,
                "rule": ind.rule,
                "purpose": ind.purpose,
                "note": ind.note,
                "loads": {
                    t.id: {
                        "tickets": ind.loads[t.id].tickets,
                        "total_minutes": ind.loads[t.id].total_minutes,
                        "weighted_load": ind.loads[t.id].weighted_load,
                    }
                    for t in sim.technicians
                },
                "max_minutes": ind.max_minutes,
                "min_minutes": ind.min_minutes,
                "gap_minutes": ind.gap_minutes,
                "balance_index": ind.balance_index,
                "stdev_minutes": ind.stdev_minutes,
                "jains_index": ind.jains_index,
                "gap_weighted": ind.gap_weighted,
                "jains_index_weighted": ind.jains_index_weighted,
                "avg_waiting": ind.avg_waiting,
                "max_waiting": ind.max_waiting,
                "last_finish": ind.last_finish.isoformat() if ind.last_finish else None,
                "total_tickets": ind.total_tickets,
                "unique_cases": ind.unique_cases,
                "total_duration": ind.total_duration,
                "ideal_load": ind.ideal_load,
                "assignments": [
                    {
                        "ticket_id": a.ticket.id,
                        "ticket_type": a.ticket.ticket_type.value,
                        "timestamp": a.arrival.isoformat(),
                        "category": a.ticket.category,
                        "priority": a.ticket.priority.value,
                        "duration": a.ticket.duration,
                        "technician": a.tech_id,
                        "status": a.status.value,
                        "start": a.start.isoformat(),
                        "finish": a.finish.isoformat(),
                        "waiting": a.waiting,
                        "rule": a.rule_note,
                        "weighted_load": a.weighted_load,
                    }
                    for a in ind.assignments
                ],
            }
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return 0

    if args.format in ("csv", "xlsx"):
        base = args.out or f"hasil_ticketing.{args.format}"
        if args.format == "csv":
            path = export_csv(indicators, sim.technicians, base)
            det = os.path.splitext(base)[0] + "_detail.csv"
            export_assignments_csv(
                indicators[keys[0]].assignments, det
            )
            print(f"Diekspor: {path}")
            print(f"Diekspor: {det}")
        else:
            path = export_xlsx(indicators, sim.technicians, base)
            print(f"Diekspor: {path}")
        return 0

    # format teks (default)
    print(text_report(indicators, sim.technicians))
    if args.detail:
        for k in keys:
            print(f"\n=== Detail Assignment {k} ===")
            print(assignment_table(indicators[k].assignments))
    return 0


def cmd_tickets_list(args) -> int:
    ds = load_dataset(args.data)
    print(f"{'ID':<7} {'Type':<19} {'Timestamp':<17} {'Category':<20} "
          f"{'Priority':<9} {'Dur':<5} {'Related':<9} Detail")
    print("─" * 110)
    for t in sorted(ds.tickets, key=lambda x: x.timestamp):
        print(f"{t.id:<7} {t.ticket_type.value:<19} "
              f"{t.timestamp.strftime('%Y-%m-%d %H:%M'):<17} "
              f"{t.category:<20} {t.priority.value:<9} {t.duration:<5} "
              f"{t.related_id or '-':<9} {t.detail}")
    print(f"\nTotal: {len(ds.tickets)} tiket | "
          f"{sum(1 for t in ds.tickets if t.is_root_case)} kasus unik")
    return 0


def cmd_tickets_add(args) -> int:
    ds = load_dataset(args.data)
    related = args.related
    requested = args.requested_by
    if args.type == "tech":
        if not related or not requested:
            print("ERROR: Technician Request wajib punya --related dan --requested-by",
                  file=sys.stderr)
            return 2
    t = add_ticket(
        dataset=ds,
        ticket_type="User Report" if args.type == "user" else "Technician Request",
        created_by="User" if args.type == "user" else "Admin",
        timestamp=args.time,
        category=args.category,
        detail=args.detail,
        priority=args.priority,
        duration=args.duration,
        related_id=related,
        requested_by=requested,
    )
    save_dataset(ds, args.data)
    print(f"Tiket ditambahkan: {t.id} ({t.ticket_type.value}) "
          f"[{t.category}/{t.priority.value}] {t.duration} menit — {t.detail}")
    print(f"Bobot prioritas: {t.weight} | Disimpan ke: {args.data or DEFAULT_DATA_PATH}")
    return 0


def cmd_technicians(args) -> int:
    ds = load_dataset(args.data)
    print(f"{'ID':<6} {'Nama':<15} {'Status Awal':<12}")
    print("─" * 40)
    for t in ds.technicians:
        print(f"{t.id:<6} {t.name:<15} {t.initial_status.value:<12}")
    return 0


def _to_day_fraction(value):
    """Nilai di sheet Excel bisa berupa fraksi hari, datetime, atau timedelta."""
    import datetime as _dt

    if value is None:
        return None
    if isinstance(value, _dt.timedelta):
        return value.total_seconds() / 86400.0
    if isinstance(value, _dt.datetime):
        return (value - _dt.datetime(1899, 12, 30)).total_seconds() / 86400.0
    return float(value)


def _expected_from_excel(path: str):
    """Baca nilai indikator langsung dari sheet Skenario Excel asli."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Skenario"]
    expected = {}
    for row in ws.iter_rows(min_row=2, max_row=5, values_only=True):
        key = str(row[0]).split()[0]  # "S1 Baseline" -> "S1"
        # Hanya S1-S3 yang punya nilai; S4 di Excel ditandai "—" (belum diimplementasikan)
        if key not in ("S1", "S2", "S3"):
            continue
        expected[key] = {
            "max_day": _to_day_fraction(row[3]),
            "min_day": _to_day_fraction(row[4]),
            "gap_day": _to_day_fraction(row[5]),
            "balance": _to_day_fraction(row[6]),
            "stdev_day": _to_day_fraction(row[7]),
            "jains": _to_day_fraction(row[8]),
            "gap_weighted": _to_day_fraction(row[9]),
            "jains_weighted": _to_day_fraction(row[10]),
            "avg_waiting_day": _to_day_fraction(row[11]),
            "max_waiting": _to_day_fraction(row[12]),
            "last_finish": row[13],
        }
    # beban per teknisi dari sheet Ringkasan
    ws2 = wb["Ringkasan"]
    loads = {}
    header = [str(c) if c is not None else "" for c in next(ws2.iter_rows(min_row=10, max_row=10, values_only=True))]
    for row in ws2.iter_rows(min_row=11, max_row=14, values_only=True):
        tech = str(row[0])
        loads[tech] = {}
        for i, h in enumerate(header):
            if h.startswith("S") and "Durasi" in h:
                key = h.split()[0]
                loads[tech][key] = _to_day_fraction(row[i])
    return expected, loads


def cmd_compare_excel(args) -> int:
    """Bandingkan hasil program dengan sheet Skenario + Ringkasan di Excel."""
    path = args.path or EXCEL_PATH
    if not os.path.exists(path):
        print(f"ERROR: file tidak ditemukan: {path}", file=sys.stderr)
        return 2

    sim = Simulation.load(args.data)
    # Override parameters with external config if present (MVP uses config.yaml in project root)
    try:
        from .config import load_parameters
        sim.dataset.parameters = load_parameters()
    except Exception:
        # If config missing or broken we keep JSON defaults – harmless for MVP
        pass
    expected, expected_loads = _expected_from_excel(path)
    tol = 1e-6

    failures = []
    checks = 0
    for key in ("S1", "S2", "S3"):
        if key not in expected:
            continue
        ind = sim.run(key)
        e = expected[key]

        def chk(label, got, exp):
            nonlocal checks
            checks += 1
            if exp is None:
                return
            ok = got is not None and abs(float(got) - float(exp)) <= tol
            if not ok:
                failures.append(f"{key}.{label}: program={got} excel={exp}")

        chk("max_minutes", ind.max_minutes / 1440, e["max_day"])
        chk("min_minutes", ind.min_minutes / 1440, e["min_day"])
        chk("gap_minutes", ind.gap_minutes / 1440, e["gap_day"])
        chk("balance_index", ind.balance_index, e["balance"])
        chk("stdev_minutes", ind.stdev_minutes / 1440, e["stdev_day"])
        chk("jains_index", ind.jains_index, e["jains"])
        chk("gap_weighted", float(ind.gap_weighted), e["gap_weighted"])
        chk("jains_weighted", ind.jains_index_weighted, e["jains_weighted"])
        chk("avg_waiting", ind.avg_waiting, e["avg_waiting_day"])  # dalam menit
        chk("max_waiting", float(ind.max_waiting), e["max_waiting"])

        # beban per teknisi (sheet Ringkasan, kolom "Sx Total Durasi")
        for tech in sim.technicians:
            if tech.id in expected_loads and key in expected_loads[tech.id]:
                exp = expected_loads[tech.id][key]
                checks += 1
                got = float(ind.loads[tech.id].total_minutes) / 1440
                if abs(got - float(exp)) > tol:
                    failures.append(
                        f"{key}.load.{tech.id}: program={got} excel={exp}"
                    )

    print(f"Membandingkan dengan: {path}")
    print(f"Total pengecekan    : {checks}")
    if failures:
        print(f"SELISIH ditemukan   : {len(failures)}")
        for f in failures:
            print(f"  ✗ {f}")
        return 1
    print("HASIL               : IDENTIK dengan Excel ✓ (toleransi 1e-6)")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="ticketing",
        description="Simulasi distribusi beban tiket teknisi "
                    "(replikasi dari file Excel Simulasi_Workload_Ticketing)",
    )
    p.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    p.add_argument("--data", default=None, help="path file dataset JSON "
                                                "(default: data/tickets.json)")
    sub = p.add_subparsers(dest="command", required=True)

    # run
    run = sub.add_parser("run", help="jalankan simulasi")
    run.add_argument("--scenario", choices=["all"] + list(SCENARIOS), default="all")
    run.add_argument("--format", choices=["text", "json", "csv", "xlsx"], default="text")
    run.add_argument("--out", help="path output untuk format csv/xlsx")
    run.add_argument("--detail", action="store_true",
                     help="tampilkan tabel assignment per tiket")
    run.set_defaults(func=cmd_run)

    # tickets
    tk = sub.add_parser("tickets", help="kelola tiket")
    tk_sub = tk.add_subparsers(dest="subcommand", required=True)
    tk_list = tk_sub.add_parser("list", help="daftar tiket")
    tk_list.set_defaults(func=cmd_tickets_list)
    tk_add = tk_sub.add_parser("add", help="tambah tiket")
    tk_add.add_argument("--type", choices=["user", "tech"], required=True,
                        help="user=User Report (T-), tech=Technician Request (TK-)")
    tk_add.add_argument("--time", type=_parse_time, required=True,
                        help="waktu tiket masuk 'YYYY-MM-DD HH:MM'")
    tk_add.add_argument("--category", required=True)
    tk_add.add_argument("--detail", required=True, help="detail kasus")
    tk_add.add_argument("--priority", choices=["Low", "Medium", "High"], required=True)
    tk_add.add_argument("--duration", type=int, required=True, help="durasi layanan (menit)")
    tk_add.add_argument("--related", help="Related Ticket ID (TK- wajib)")
    tk_add.add_argument("--requested-by", help="ID teknisi peminta (TK- wajib)")
    tk_add.set_defaults(func=cmd_tickets_add)

    # technicians
    tc = sub.add_parser("technicians", help="daftar teknisi")
    tc.set_defaults(func=cmd_technicians)

    # compare-excel
    ce = sub.add_parser("compare-excel", help="validasi hasil terhadap file Excel asli")
    ce.add_argument("path", nargs="?", help="path file .xlsx "
                                            "(default: path Excel asli)")
    ce.set_defaults(func=cmd_compare_excel)

    return p


def main(argv=None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
