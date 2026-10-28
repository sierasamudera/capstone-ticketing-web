"""Tabel laporan untuk CLI + ekspor CSV/XLSX.

Output tabel meniru sheet Sim_* (per tiket) dan sheet Ringkasan/Skenario
(indikator), agar mudah dibandingkan langsung dengan file Excel.
"""

from __future__ import annotations

import csv
import os
from datetime import datetime
from typing import Optional

from .metrics import BalanceIndicators, LoadSummary, scenario_ranking
from .models import Assignment

BAR = "─"
WIDTH = 118


def fmt_dt(dt: Optional[datetime]) -> str:
    return dt.strftime("%Y-%m-%d %H:%M") if dt else "—"


def fmt_hmm(minutes: float) -> str:
    """Format menit ke [h]:mm (sheet Ringkasan menampilkan '4:05:00')."""
    total = int(round(minutes))
    h, m = divmod(total, 60)
    return f"{h}:{m:02d}"


def assignment_table(assignments: list) -> str:
    """Tabel per tiket seperti sheet Sim_S1/S2/S3 (kolom A-N + P)."""
    headers = [
        "Ticket ID", "Type", "Timestamp", "Category", "Priority",
        "Dur", "Tech", "Status", "Start", "Est. Finish", "Wait",
        "Weighted", "Assignment Rule",
    ]
    rows = []
    for a in assignments:
        rows.append([
            a.ticket.id,
            "T-" if a.ticket.ticket_type.value == "User Report" else "TK-",
            a.arrival.strftime("%H:%M"),
            a.ticket.category,
            a.ticket.priority.value,
            str(a.ticket.duration),
            a.tech_id,
            a.status.value,
            a.start.strftime("%H:%M"),
            a.finish.strftime("%H:%M"),
            str(a.waiting),
            str(a.weighted_load),
            a.rule_note,
        ])
    return _table(headers, rows)


def loads_table(indicator: BalanceIndicators, technicians: list) -> str:
    """Tabel per teknisi seperti sheet Ringkasan R10-R14."""
    headers = ["Teknisi", "Jumlah Tiket", "Total Durasi", "Beban Berbobot (pt-min)"]
    rows = []
    for t in technicians:
        load: LoadSummary = indicator.loads[t.id]
        rows.append([
            f"{t.id} ({t.name})",
            str(load.tickets),
            fmt_hmm(load.total_minutes),
            str(load.weighted_load),
        ])
    return _table(headers, rows)


def indicator_table(indicators: dict) -> str:
    """Tabel indikator seperti sheet Skenario R2-R5."""
    headers = [
        "Skenario", "Max", "Min", "Gap", "Min/Max", "Std Dev",
        "Jain's", "Gap W", "Jain's W", "Avg Wait", "Max Wait", "Selesai",
    ]
    rows = []
    for key, ind in indicators.items():
        rows.append([
            ind.name,
            fmt_hmm(ind.max_minutes),
            fmt_hmm(ind.min_minutes),
            fmt_hmm(ind.gap_minutes),
            f"{ind.balance_index:.4f}",
            f"{ind.stdev_minutes:.2f}",
            f"{ind.jains_index:.4f}",
            str(int(ind.gap_weighted)),
            f"{ind.jains_index_weighted:.4f}",
            f"{ind.avg_waiting:.2f}",
            str(int(ind.max_waiting)),
            fmt_dt(ind.last_finish),
        ])
    return _table(headers, rows)


def _table(headers: list, rows: list) -> str:
    """Renderer tabel sederhana dengan auto-width."""
    all_rows = [headers] + rows
    widths = [max(len(str(r[i])) for r in all_rows) for i in range(len(headers))]

    def sep() -> str:
        return BAR + BAR.join(BAR * (w + 2) for w in widths) + BAR

    def render(row: list) -> str:
        return " " + " | ".join(str(cell).ljust(w) for cell, w in zip(row, widths)) + " "

    lines = [sep(), render(headers), sep()]
    for r in rows:
        lines.append(render(r))
    lines.append(sep())
    return "\n".join(lines)


def ranking_text(indicators: dict) -> str:
    """Kesimpulan seperti sheet Skenario R7-R10."""
    r = scenario_ranking(indicators)
    if not r:
        return "(tidak ada skenario untuk diperingkat)"
    lines = []
    name = {k: v.name for k, v in indicators.items()}
    if "best_gap_minutes" in r:
        lines.append(f"Skenario dengan Gap menit terkecil   : {name[r['best_gap_minutes']]}")
    if "best_gap_weighted" in r:
        lines.append(f"Skenario dengan Gap weighted terkecil : {name[r['best_gap_weighted']]}")
    if "reduction_S2" in r:
        lines.append(f"Penurunan Gap menit S2 vs S1          : {r['reduction_S2']:.1%}")
    if "reduction_S3" in r:
        lines.append(f"Penurunan Gap weighted S3 vs S1       : {r['reduction_S3']:.1%}")
    return "\n".join(lines)


def text_report(indicators: dict, technicians: list, scenario_key: Optional[str] = None) -> str:
    """Laporan teks lengkap untuk satu atau semua skenario."""
    parts = []
    title = "SIMULASI DISTRIBUSI BEBAN TIKET"
    parts.append(title)
    parts.append(BAR * WIDTH)

    first = indicators[scenario_key] if scenario_key else next(iter(indicators.values()))
    parts.append(f"Total tiket        : {first.total_tickets}")
    parts.append(f"Jumlah kasus unik  : {first.unique_cases}")
    parts.append(f"Total durasi       : {first.total_duration} menit")
    parts.append(f"Beban ideal/teknisi: {first.ideal_load:.2f} menit")
    parts.append("")

    keys = [scenario_key] if scenario_key else list(indicators)
    for key in keys:
        ind = indicators[key]
        parts.append(f"=== {ind.name} — {ind.rule} ===")
        if not ind.implemented:
            parts.append(f"  {ind.note}")
        else:
            parts.append(loads_table(ind, technicians))
            parts.append("")
            parts.append(f"  Gap (menit)         : {ind.gap_minutes:.0f}  "
                         f"[Excel: {ind.gap_minutes_day:.6f}]")
            parts.append(f"  Min/Max             : {ind.balance_index:.6f}")
            parts.append(f"  Std Dev (menit)     : {ind.stdev_minutes:.6f}")
            parts.append(f"  Jain's Index        : {ind.jains_index:.6f}")
            parts.append(f"  Gap Weighted (pt-m) : {ind.gap_weighted:.0f}")
            parts.append(f"  Jain's Weighted     : {ind.jains_index_weighted:.6f}")
            parts.append(f"  Waiting avg/max     : {ind.avg_waiting:.2f} / {ind.max_waiting:.0f} menit")
            parts.append(f"  Selesai terakhir    : {fmt_dt(ind.last_finish)}")
            parts.append(f"  Catatan             : {ind.note}")
        parts.append("")

    if not scenario_key:
        parts.append("=== Perbandingan Indikator ===")
        parts.append(indicator_table(indicators))
        parts.append("")
        parts.append(ranking_text(indicators))

    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Ekspor
# ---------------------------------------------------------------------------


def export_csv(indicators: dict, technicians: list, path: str) -> str:
    """Ekspor hasil semua skenario ke CSV."""
    rows = []
    for key, ind in indicators.items():
        for t in technicians:
            load = ind.loads[t.id]
            rows.append([
                key, ind.name, t.id, t.name, load.tickets,
                load.total_minutes, load.weighted_load,
                ind.gap_minutes, ind.balance_index, ind.stdev_minutes,
                ind.jains_index, ind.gap_weighted, ind.jains_index_weighted,
                ind.avg_waiting, ind.max_waiting,
                ind.last_finish.strftime("%Y-%m-%d %H:%M") if ind.last_finish else "",
            ])
    headers = [
        "Skenario", "Nama", "Tech ID", "Nama Teknisi", "Jumlah Tiket",
        "Total Durasi (min)", "Weighted (pt-min)", "Gap (min)", "Min/Max",
        "Std Dev (min)", "Jain's", "Gap Weighted", "Jain's Weighted",
        "Avg Waiting", "Max Waiting", "Selesai Terakhir",
    ]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)
    return path


def export_assignments_csv(assignments: list, path: str) -> str:
    """Ekspor detail assignment per tiket ke CSV (seperti sheet Sim_*)."""
    headers = [
        "Ticket ID", "Ticket Type", "Timestamp", "Category", "Detail Case",
        "Priority", "Priority Weight", "Duration (min)", "Assigned Technician",
        "Tech Status at Arrival", "Start Time", "Estimated Finish",
        "Waiting Time (min)", "Assignment Rule", "Weighted Load (pt-min)",
        "Related Ticket ID", "Requested By (Technician)",
    ]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for a in assignments:
            w.writerow([
                a.ticket.id, a.ticket.ticket_type.value,
                a.ticket.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                a.ticket.category, a.ticket.detail, a.ticket.priority.value,
                a.ticket.weight, a.ticket.duration, a.tech_id, a.status.value,
                a.start.strftime("%Y-%m-%d %H:%M:%S"),
                a.finish.strftime("%Y-%m-%d %H:%M:%S"),
                a.waiting, a.rule_note, a.weighted_load,
                a.ticket.related_id or "", a.ticket.requested_by or "",
            ])
    return path


def export_xlsx(indicators: dict, technicians: list, path: str) -> str:
    """Ekspor hasil ke workbook Excel (butuh openpyxl)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # --- sheet Ringkasan ---
    ws = wb.active
    ws.title = "Ringkasan"
    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="DDE6F0")

    ws.append(["Indikator", "Nilai", "Keterangan"])
    for c in ws[1]:
        c.font = bold
        c.fill = hdr_fill
    first = next(iter(indicators.values()))
    ws.append(["Jumlah tiket (total baris)", first.total_tickets,
               "Termasuk tiket User Report dan Technician Request"])
    ws.append(["Jumlah kasus (unik)", first.unique_cases,
               "Tiket tanpa Related Ticket ID = kasus baru (root)"])
    ws.append(["Total service duration (min)", first.total_duration,
               "Total estimasi waktu layanan, semua jenis tiket"])
    ws.append(["Beban ideal per teknisi (min)", round(first.ideal_load, 2),
               "Total durasi ÷ jumlah teknisi"])
    ws.append([])
    ws.append(["Teknisi"] + sum(
        [[f"{k} Jumlah", f"{k} Durasi", f"{k} Weighted"] for k in indicators], []
    ))
    for c in ws[ws.max_row]:
        c.font = bold
        c.fill = hdr_fill
    for t in technicians:
        row = [f"{t.id} ({t.name})"]
        for key, ind in indicators.items():
            load = ind.loads[t.id]
            row += [load.tickets, load.total_minutes, load.weighted_load]
        ws.append(row)

    # --- sheet per skenario (detail assignment) ---
    for key, ind in indicators.items():
        ws2 = wb.create_sheet(title=f"Sim_{key}"[:31])
        headers = [
            "Ticket ID", "Type", "Timestamp", "Category", "Detail Case",
            "Priority", "Weight", "Duration (min)", "Technician",
            "Status at Arrival", "Start Time", "Estimated Finish",
            "Waiting (min)", "Assignment Rule", "Weighted Load (pt-min)",
        ]
        ws2.append(headers)
        for c in ws2[1]:
            c.font = bold
            c.fill = hdr_fill
        for a in ind.assignments:
            ws2.append([
                a.ticket.id, a.ticket.ticket_type.value,
                a.ticket.timestamp.strftime("%Y-%m-%d %H:%M:%S"),
                a.ticket.category, a.ticket.detail, a.ticket.priority.value,
                a.ticket.weight, a.ticket.duration, a.tech_id, a.status.value,
                a.start.strftime("%Y-%m-%d %H:%M:%S"),
                a.finish.strftime("%Y-%m-%d %H:%M:%S"),
                a.waiting, a.rule_note, a.weighted_load,
            ])
        for col in range(1, len(headers) + 1):
            ws2.column_dimensions[get_column_letter(col)].width = 22

    for col in ("A", "B", "C"):
        ws.column_dimensions[col].width = 42
    wb.save(path)
    return path
